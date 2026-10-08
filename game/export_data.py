#!/usr/bin/env python3
"""Xuất số liệu từ file Excel (Goma_Game_Data_v9.xlsx) ra data.json cho game đọc.
Chạy lại mỗi khi anh sửa Excel:  python export_data.py [đường_dẫn_xlsx]
Lưu ý: cần mở Excel rồi LƯU (hoặc chạy recalc) để các ô công thức có giá trị."""
import sys, json, warnings; warnings.filterwarnings("ignore")
from openpyxl import load_workbook
import os
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "data", "Goma_Game_Data_v9.xlsx")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "data.json")
wb = load_workbook(SRC, data_only=True)
V = lambda ws, a: wb[ws][a].value
num = lambda x: None if x in (None, "") else x

# --- tướng ---
heroes = {}
n = wb["NHÂN_VẬT"]
for r in range(6, 35):
    code, name = n.cell(row=r, column=1).value, n.cell(row=r, column=2).value
    if not code or not name: continue
    heroes[code] = dict(code=code, name=str(name).strip(), role=n.cell(row=r, column=4).value,
        row="front" if n.cell(row=r, column=5).value == "Hàng trước" else "back",
        bio=n.cell(row=r, column=6).value, quote=n.cell(row=r, column=7).value, vip=('VIP' in str(n.cell(row=r, column=8).value or '')), tiers=[], skill=dict(components=[]), passive=dict(components=[]))
c = wb["CHỈ_SỐ_TÍNH"]; TIERS = ["Trắng", "Xanh lá", "Xanh dương"]
for r in range(5, 160):
    code, tier = c.cell(row=r, column=1).value, c.cell(row=r, column=3).value
    if code in heroes and tier in TIERS:
        g = lambda col: c.cell(row=r, column=col).value
        heroes[code]["tiers"].append(dict(tier=tier, hp=g(4), atk=g(5), df=g(6), spd=g(7), regen=g(8), dodge=g(9), acc=g(10), crit=g(11), critDmg=g(12), critRes=g(13)))
d = wb["SKILL_MÔ_TẢ"]
for r in range(5, 70):
    code, kind = d.cell(row=r, column=1).value, d.cell(row=r, column=3).value
    if code in heroes:
        tgt = heroes[code]["skill" if kind == "Skill nộ" else "passive"]
        tgt["name"] = d.cell(row=r, column=4).value
        tgt["desc"] = [d.cell(row=r, column=x).value for x in (5, 6, 7)]
t = wb["SKILL_THÀNH_PHẦN"]; FROM = {"Trắng": 0, "Xanh lá": 1, "Xanh dương": 2}
for r in range(5, 85):
    code = t.cell(row=r, column=1).value
    if code not in heroes: continue
    g = lambda col: t.cell(row=r, column=col).value
    comp = dict(kind=g(6), target=g(7), trigger=g(4), frm=FROM.get(g(5), 0), unit=g(11),
        n=[num(g(8)), num(g(9)), num(g(10))], mag=[num(g(12)), num(g(13)), num(g(14))],
        prob=[num(g(15)), num(g(16)), num(g(17))], dur=num(g(18)), status=g(19), note=g(23))
    heroes[code]["skill" if g(3) == "Skill nộ" else "passive"]["components"].append(comp)
# --- trạng thái ---
fx = {}
h = wb["HIỆU_ỨNG"]
for r in range(5, 45):
    nm = h.cell(row=r, column=1).value
    if nm: fx[nm] = dict(name=nm, kind=h.cell(row=r, column=2).value, desc=h.cell(row=r, column=4).value, dur=num(h.cell(row=r, column=5).value))
data = dict(heroes=list(heroes.values()), statuses=fx, gacha=dict(tiers=TIERS, weights=[65, 27, 8]),
            source=SRC.split("/")[-1])

# --- quái, boss, map ---
mon = {}
q = wb["QUÁI_BOSS"]
for r in range(5, 70):
    code = q.cell(row=r, column=1).value
    if not code: continue
    g = lambda c: q.cell(row=r, column=c).value
    mon[code] = dict(code=code, name=g(2), kind=g(3), map=str(g(4)), row="front" if g(5) == "Hàng trước" else "back",
        stats=dict(hp=g(6), atk=g(7), df=g(8), spd=g(9), regen=g(10) or 0, dodge=g(11) or 0, acc=g(12) or 0, crit=g(13) or 0, critDmg=g(14) or 1.5, critRes=g(15) or 0),
        skill=dict(name=g(16), desc=g(17), components=[]), passive=dict(name=g(18), desc=g(19), components=[]), note=g(20))
qc = wb["QUÁI_THÀNH_PHẦN"]
for r in range(5, 140):
    code = qc.cell(row=r, column=1).value
    if code not in mon: continue
    g = lambda c: qc.cell(row=r, column=c).value
    n, mag, prob = g(7), g(9), g(10)
    comp = dict(kind=g(5), target=g(6), trigger=g(4), frm=0, unit=g(8), n=[n, n, n], mag=[mag, mag, mag], prob=[prob, prob, prob], dur=g(11), status=g(12), note=g(13))
    mon[code]["skill" if g(3) == "Skill nộ" else "passive"]["components"].append(comp)
levels = []
mm = wb["MAP_MÀN"]
for r in range(5, 100):
    mp = mm.cell(row=r, column=1).value
    if mp is None: continue
    g = lambda c: mm.cell(row=r, column=c).value
    comp = []
    for tok in str(g(5)).split(";"):
        tok = tok.strip()
        if not tok: continue
        code, cnt = tok.split("×"); comp.append(dict(code=code.strip(), count=int(cnt)))
    levels.append(dict(map=int(mp), man=int(g(2)), name=g(3), place=g(4), comp=comp, kind=g(6), hpMul=g(7), atkMul=g(8), team=g(9), win=g(10), turns=g(11), reward=g(12)))
data["monsters"] = list(mon.values()); data["levels"] = levels

json.dump(data, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
open(OUT.replace(".json", ".js"), "w", encoding="utf-8").write("window.GOMA_DATA = " + json.dumps(data, ensure_ascii=False) + ";\n")   # để index.html mở thẳng bằng file:// vẫn đọc được
print(f"{len(mon)} quái, {len(levels)} màn, ", end=''); print(f"{len(heroes)} tướng, {sum(len(x['skill']['components'])+len(x['passive']['components']) for x in heroes.values())} thành phần, {len(fx)} hiệu ứng -> {OUT}")
for x in heroes.values(): print(" ", x["code"], x["name"], "|", x["role"], "|", len(x["tiers"]), "bậc |", len(x["skill"]["components"]), "skill |", len(x["passive"]["components"]), "nội tại")
