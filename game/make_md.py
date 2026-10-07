#!/usr/bin/env python3
"""Tạo lại file md thiết kế từ Excel. Chạy trong thư mục game/:  python make_md.py"""
import os, warnings; warnings.filterwarnings("ignore")
from itertools import groupby
from collections import Counter
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "..", "data", "Goma_Game_Data_v8.xlsx"); OUT = os.path.join(HERE, "..", "docs", "Goma_Game_Thiet_Ke_v6.md")
wb = load_workbook(XLSX, data_only=True)
cell = lambda x: "" if x is None else str(x).replace("\n", " ").replace("|", "/")
def table(h, rows): return "\n".join(["| " + " | ".join(h) + " |", "|" + "|".join("---" for _ in h) + "|"] + ["| " + " | ".join(cell(x) for x in r) + " |" for r in rows])
L = wb["LUẬT_VÀ_QUYẾT_ĐỊNH"]
dec = sorted([tuple((L.cell(row=r, column=c).value or "") for c in range(1, 6)) for r in range(5, L.max_row + 1) if L.cell(row=r, column=1).value], key=lambda d: d[0])
cnt = Counter(d[3] for d in dec)
o = ["# Game thẻ bài GOMA – Tổng hợp thiết kế (v6)\n", "Tạo tự động từ `data/Goma_Game_Data_v8.xlsx` bằng `game/make_md.py`.\n", "## Đọc nhanh\n",
 "* (*) Bản thô chơi được: **10 tướng** × 3 phẩm chất, Thủ kho Hiếu có sẵn, lần rút đầu ra tướng Trắng khác Hiếu. Không có mini-game ở đợt này. 3 map × 10 màn (map 1 đã hiệu chỉnh, map 2 và 3 chưa).",
 f"* (*) Trạng thái quyết định: {cnt['Đã chốt']} đã chốt, {cnt['Tạm']} tạm, {cnt['Chưa chốt']} chưa chốt, {cnt['Treo']} đang treo.",
 "* (*) Mọi con số (chỉ số, skill, quái, hệ số màn) là SỐ NHÁP để hỗ trợ, chưa cân. Cân bằng thật dựa vào chơi thử và núm chỉnh toàn cục của engine.\n"]
for g, items in groupby(dec, key=lambda d: d[0]):
    o += [f"## {g}\n", table(["Mục", "Nội dung", "Trạng thái", "Ghi chú"], [(i[1], i[2], i[3], i[4]) for i in items]), ""]
n, b, d_ = wb["NHÂN_VẬT"], wb["CHỈ_SỐ_TRẮNG"], wb["SKILL_MÔ_TẢ"]
names = {}
for r in range(5, 70):
    c = d_.cell(row=r, column=1).value
    if c: names.setdefault(c, {})[d_.cell(row=r, column=3).value] = d_.cell(row=r, column=4).value
rows = []
for r in range(6, 35):
    if n.cell(row=r, column=2).value:
        c = n.cell(row=r, column=1).value; nm = names.get(c, {})
        rows.append((c, n.cell(row=r, column=2).value, n.cell(row=r, column=4).value, n.cell(row=r, column=5).value, *[b.cell(row=r, column=x).value for x in (4, 5, 6, 7, 8)], nm.get("Skill nộ") or "", nm.get("Nội tại") or ""))
o += ["## Mười tướng (chỉ số bản Trắng; tên skill là TẠM)\n", table(["Mã", "Tên", "Vai", "Hàng", "Máu", "Công", "Thủ", "Tốc độ", "Hồi/lượt", "Skill nộ", "Nội tại"], rows), ""]
k = wb["KHUNG_VAI_TRÒ"]
o += ["## Khung vai trò (hệ số so với chuẩn Trắng: máu 1000, công 100, thủ 40, tốc độ 100)\n", table(["Vai trò", "Máu", "Công", "Thủ", "Tốc độ", "Kiểu skill", "Ngân sách"], [(k.cell(row=r, column=1).value, f"{k.cell(row=r,column=2).value}-{k.cell(row=r,column=3).value}", f"{k.cell(row=r,column=4).value}-{k.cell(row=r,column=5).value}", f"{k.cell(row=r,column=6).value}-{k.cell(row=r,column=7).value}", f"{k.cell(row=r,column=8).value}-{k.cell(row=r,column=9).value}", k.cell(row=r, column=10).value, k.cell(row=r, column=11).value) for r in range(9, 16)]), ""]
q = wb["QUÁI_BOSS"]
o += ["## Quái và boss (BẢN NHÁP)\n", table(["Mã", "Quái", "Loại", "Map", "Hàng", "Máu", "Công", "Thủ", "Tốc độ", "Skill", "Nội tại"], [tuple(q.cell(row=r, column=c).value for c in (1, 2, 3, 4, 5, 6, 7, 8, 9, 17, 19)) for r in range(5, 25) if q.cell(row=r, column=1).value]), ""]
m = wb["MAP_MÀN"]
o += ["## 30 màn (hệ số map 1 do mô phỏng; map 2, 3 chưa hiệu chỉnh)\n", table(["Map", "Màn", "Tên", "Quái", "Loại", "Hệ số máu", "Hệ số công", "Đội", "Thắng MP", "Lượt"], [tuple(m.cell(row=r, column=c).value for c in (1, 2, 3, 5, 6, 7, 8, 9, 10, 11)) for r in range(5, 35) if m.cell(row=r, column=1).value is not None]), ""]
o += ["## Số đo mô phỏng engine\n", "```", open(os.path.join(HERE, "sim_report.txt"), encoding="utf-8").read().strip(), "```\n"]
open(OUT, "w", encoding="utf-8").write("\n".join(o)); print("md ok:", OUT)
